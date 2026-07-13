from fastapi import APIRouter, HTTPException, File, UploadFile, Depends
import logging
import csv
import io
import re
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List
from database import get_database
from authentication_service import get_current_user
from rebac_service import get_reback_service, ReBACService

logger = logging.getLogger(__name__)

router = APIRouter()
rebac: ReBACService = get_reback_service()

_SUPER_ROLES = {"Super Admin", "super_admin", "admin", "platform-admin"}


def _now_date() -> str:
    return datetime.now(timezone.utc).date().isoformat()


def _is_super(role: str) -> bool:
    return role in _SUPER_ROLES


def _user_role(user) -> str:
    if isinstance(user, dict):
        return user.get("role") or ""
    return getattr(user, "role", None) or ""


def _user_tenant(user) -> str:
    if isinstance(user, dict):
        return user.get("tenant_id") or user.get("tenantId") or ""
    return getattr(user, "tenant_id", None) or ""


async def _resolve_fw_filter(
    db, framework_id: str, tid: str, role: str
) -> Dict[str, Any]:
    """
    Return a MongoDB filter for a mutation on a framework.

    Super admins can write to any framework (including global seeded ones).
    Tenant admins may only write to frameworks they own (tenant_id == tid).
    If a tenant admin targets a global framework we raise 403 with a helpful
    message — they should clone it first via POST /api/compliance.
    """
    if _is_super(role):
        return {"id": framework_id}

    # Tenant-scoped write
    own_fw = await db.compliance_frameworks.find_one(
        {"id": framework_id, "tenant_id": tid}, {"_id": 1}
    )
    if own_fw:
        return {"id": framework_id, "tenant_id": tid}

    # Check whether the framework exists at all (as a global one)
    exists = await db.compliance_frameworks.find_one({"id": framework_id}, {"_id": 1})
    if exists:
        raise HTTPException(
            status_code=403,
            detail=(
                "This is a shared framework — modifying it would affect all tenants. "
                "Use 'Add Compliance → Standard Frameworks' to create your own copy, "
                "then add controls to that copy."
            ),
        )
    raise HTTPException(status_code=404, detail="Framework not found")


def _parse_controls_from_text(text: str) -> List[Dict[str, Any]]:
    """
    Heuristic extraction of controls from plain text (PDF / DOCX).
    Matches lines that start with an alphanumeric control ID, e.g.
    AC-1, CC1.1, A.6.1, 1.2.3, ISO-5.1.
    """
    controls: List[Dict[str, Any]] = []
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    id_pat = re.compile(r'^([A-Z]{1,5}[-.]?\d+(?:[-.]\d+)*)\s{1,3}(.+)', re.IGNORECASE)
    num_pat = re.compile(r'^(\d+\.\d+(?:\.\d+)?)\s{1,3}(.+)')

    i = 0
    while i < len(lines):
        m = id_pat.match(lines[i]) or num_pat.match(lines[i])
        if m:
            ctrl_id, ctrl_name = m.group(1), m.group(2)[:200]
            desc_parts: List[str] = []
            j = i + 1
            while j < len(lines) and j < i + 5:
                if id_pat.match(lines[j]) or num_pat.match(lines[j]):
                    break
                desc_parts.append(lines[j])
                j += 1
            controls.append({
                "id":          ctrl_id,
                "name":        ctrl_name,
                "description": " ".join(desc_parts)[:500],
                "category":    "Imported",
                "status":      "Not Implemented",
                "lastReviewed": _now_date(),
                "evidence":    [],
            })
            i = j
        else:
            i += 1
    return controls


@router.post("/api/compliance")
async def create_compliance_framework(
    body: Dict[str, Any],
    current_user=Depends(get_current_user),
):
    """Create a new (custom or catalog) compliance framework for this tenant."""
    db = get_database()
    name = (body.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Framework name is required")

    framework_id = body.get("id") or re.sub(r"[^a-z0-9_]", "_", name.lower())[:40]
    from tenant_context import get_tenant_id
    tenant_id = get_tenant_id() or ""

    if await db.compliance_frameworks.find_one({"id": framework_id, "tenant_id": tenant_id}):
        raise HTTPException(status_code=409, detail="A framework with this ID already exists")

    framework = {
        "id":          framework_id,
        "name":        name,
        "shortName":   (body.get("shortName") or name[:20]).strip(),
        "description": (body.get("description") or "").strip(),
        "category":    (body.get("category") or "Custom").strip(),
        "official_url": (body.get("official_url") or "").strip(),
        "status":      "Pending",
        "progress":    0,
        "controls":    [],
        "tenant_id":   tenant_id,
        "created_by":  current_user.username or "",
        "created_at":  datetime.now(timezone.utc).isoformat(),
    }
    await db.compliance_frameworks.insert_one({**framework})
    framework.pop("_id", None)
    logger.info("Compliance framework created: %s by %s", framework_id, current_user.username)
    return framework


@router.get("/api/compliance")
async def get_compliance_frameworks(
    current_user=Depends(get_current_user),
):
    """
    Fetch compliance frameworks visible to the caller:
      - Global/seeded frameworks (no tenant_id or empty tenant_id)
      - Frameworks owned by the caller's tenant
    """
    db = get_database()
    from tenant_context import get_tenant_id
    tid = get_tenant_id() or _user_tenant(current_user)
    logger.debug("get_compliance_frameworks - tenant_id: %s", tid)

    # Show global frameworks (no tenant stamp) PLUS this tenant's own frameworks
    query = {"$or": [
        {"tenant_id": {"$exists": False}},
        {"tenant_id": ""},
        {"tenant_id": None},
        {"tenant_id": tid},
    ]}
    frameworks = await db.compliance_frameworks.find(query, {"_id": 0}).to_list(length=200)
    frameworks.sort(key=lambda x: x.get("name", ""))

    if not frameworks:
        logger.warning("No compliance frameworks found. Run: python backend/seed_compliance.py")

    return frameworks


@router.post("/api/compliance/{framework_id}/controls")
async def add_compliance_control(
    framework_id: str,
    control: dict,
    current_user=Depends(get_current_user),
):
    """Add a new control to a framework."""
    db = get_database()

    if not control.get("id") or not control.get("name"):
        raise HTTPException(status_code=400, detail="Control ID and Name are required")

    new_control = {
        "id": control["id"],
        "name": control["name"],
        "description": control.get("description", ""),
        "category": control.get("category", "General"),
        "status": "Not Implemented",
        "lastReviewed": datetime.now(timezone.utc).date().isoformat(),
        "evidence": [],
    }

    from tenant_context import get_tenant_id
    tid = get_tenant_id() or _user_tenant(current_user)
    fw_filter = await _resolve_fw_filter(db, framework_id, tid, _user_role(current_user))

    # ReBAC check: does the current user have 'add_control' permission on this framework?
    user_fga = f"user:{current_user.username}"
    # OpenFGA object format: type:id, e.g., "framework:framework_id"
    if not await rebac.check_permission(user_fga, "add_control", "framework", framework_id):
        raise HTTPException(status_code=403, detail="Not authorized to add controls to this framework via ReBAC")

    result = await db.compliance_frameworks.update_one(
        fw_filter,
        {"$push": {"controls": new_control}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Framework not found")

    return new_control


@router.post("/api/compliance/{framework_id}/import")
async def import_compliance_controls(
    framework_id: str,
    file: UploadFile = File(...),
    current_user=Depends(get_current_user),
):
    """
    Import controls from CSV, Excel (.xlsx), PDF, or Word (.docx).
    CSV/XLSX columns: ID, Name, Description, Category.
    PDF/DOCX: heuristic extraction by control-ID patterns.
    """
    db = get_database()
    filename = (file.filename or "").lower()
    content = await file.read()
    new_controls: List[Dict[str, Any]] = []

    # ── CSV ──────────────────────────────────────────────────────────────────
    if filename.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8", errors="replace")))
        for row in reader:
            cid  = (row.get("ID") or row.get("id") or "").strip()
            name = (row.get("Name") or row.get("name") or "").strip()
            desc = (row.get("Description") or row.get("description") or "").strip()
            cat  = (row.get("Category") or row.get("category") or "Imported").strip()
            if cid and name:
                new_controls.append({"id": cid, "name": name, "description": desc,
                                     "category": cat, "status": "Not Implemented",
                                     "lastReviewed": _now_date(), "evidence": []})

    # ── Excel (.xlsx) ─────────────────────────────────────────────────────────
    elif filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise HTTPException(status_code=400, detail="Empty spreadsheet")
            headers = [str(h).strip().lower() if h else "" for h in rows[0]]
            for row in rows[1:]:
                r = {headers[i]: (str(row[i]).strip() if row[i] is not None else "")
                     for i in range(min(len(headers), len(row)))}
                cid  = r.get("id") or r.get("control id") or ""
                name = r.get("name") or r.get("control name") or ""
                desc = r.get("description") or ""
                cat  = r.get("category") or "Imported"
                if cid and name:
                    new_controls.append({"id": cid, "name": name, "description": desc,
                                         "category": cat, "status": "Not Implemented",
                                         "lastReviewed": _now_date(), "evidence": []})
        except ImportError:
            raise HTTPException(status_code=422,
                detail="openpyxl is required for Excel import. Install it or upload a CSV.")

    # ── PDF ───────────────────────────────────────────────────────────────────
    elif filename.endswith(".pdf"):
        try:
            import pdfplumber
            text_parts: List[str] = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for page in pdf.pages:
                    text_parts.append(page.extract_text() or "")
            new_controls = _parse_controls_from_text("\n".join(text_parts))
        except ImportError:
            try:
                import PyPDF2
                reader_pdf = PyPDF2.PdfReader(io.BytesIO(content))
                text_parts = [page.extract_text() or "" for page in reader_pdf.pages]
                new_controls = _parse_controls_from_text("\n".join(text_parts))
            except ImportError:
                raise HTTPException(status_code=422,
                    detail="pdfplumber or PyPDF2 is required for PDF import.")

    # ── Word (.docx) ──────────────────────────────────────────────────────────
    elif filename.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            full_text = "\n".join(p.text for p in doc.paragraphs)
            new_controls = _parse_controls_from_text(full_text)
        except ImportError:
            raise HTTPException(status_code=422,
                detail="python-docx is required for Word import. Install it or upload a CSV.")
    else:
        raise HTTPException(status_code=415,
            detail="Unsupported file type. Upload CSV, XLSX, PDF, or DOCX.")

    if not new_controls:
        raise HTTPException(status_code=400,
            detail="No controls could be extracted. Check the file format and content.")

    from tenant_context import get_tenant_id
    tid = get_tenant_id() or _user_tenant(current_user)
    fw_filter = await _resolve_fw_filter(db, framework_id, tid, _user_role(current_user))
    result = await db.compliance_frameworks.update_one(
        fw_filter,
        {"$push": {"controls": {"$each": new_controls}}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Framework not found")

    logger.info("Imported %d controls into framework %s", len(new_controls), framework_id)
    return {"status": "success", "count": len(new_controls)}
